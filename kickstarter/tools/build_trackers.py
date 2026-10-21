"""Builds the two tracking workbooks:
  assets/data/REEFRUN_validation_tracker.xlsx   pre-campaign demand validation (02)
  11_campaign_dashboard.xlsx                     live campaign dashboard, forecasts and triggers (11)
Run from kickstarter/: python tools/build_trackers.py ; then python tools/check_workbook.py <file>
Blue = input you type in · yellow = assumption to revisit · black = formula · first data row of each log is an
EXAMPLE row (grey italic) showing the format; overwrite or delete it.
"""
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

F = "Arial"
BLUE = Font(name=F, color="0000FF")
EX = Font(name=F, color="808080", italic=True)
BLACK = Font(name=F)
BOLD = Font(name=F, bold=True)
TITLE = Font(name=F, bold=True, size=14)
HDR = Font(name=F, bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="0B3B52")
KEY = PatternFill("solid", fgColor="FFFF00")
OUTF = PatternFill("solid", fgColor="DDEFF5")
GOOD = PatternFill("solid", fgColor="C6EFCE")
BAD = PatternFill("solid", fgColor="FFC7CE")
USD = '€#,##0;(€#,##0);"-"'  # campaign currency is EUR (NL creator)
PCT = '0.0%;(0.0%);"-"'
NUM = '#,##0;(#,##0);"-"'
DATE = "yyyy-mm-dd"


def header(ws, row, labels, widths=None):
    for i, lab in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font, c.fill = HDR, HFILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 42
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def put(ws, addr, v, font=BLACK, fmt=None, fill=None, comment=None):
    c = ws[addr]
    c.value, c.font = v, font
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if comment:
        c.comment = Comment(comment, "REEF RUN")
    return c


def readme(wb, lines):
    ws = wb.active
    ws.title = "README"
    ws.column_dimensions["A"].width = 120
    for i, (t, f) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=t)
        c.font = f
        c.alignment = Alignment(wrap_text=True)


# =====================================================================================================
# 1. Validation tracker
# =====================================================================================================
wb = Workbook()
readme(wb, [
    ("REEF RUN — demand validation tracker (02_DEMAND_VALIDATION.md)", TITLE),
    ("Blue = you type it in. Black = formula. Grey italic row = an EXAMPLE showing the format; overwrite it. Yellow = threshold you may change.", BLACK),
    ("Weekly KPIs: one row per week, numbers from Steamworks (UTM + wishlist reports), the email tool, Discord and playtest logs.", BLACK),
    ("Clips: one row per short video per platform, metrics at 72 hours. Concepts: A = 'It can hear you', B = 'Be the shark', C = 'Hide in plain sight'.", BLACK),
    ("Survey: one row per playtest respondent. Only respondents with 3+ rounds count toward the purchase-intent gate.", BLACK),
    ("Gates: computed from the other sheets. Vanity numbers (views, likes, followers) are logged but never decide a gate.", BOLD),
])

ws = wb.create_sheet("Weekly KPIs")
cols = ["Week #", "Week start", "Steam page visits", "Wishlist adds", "Wishlist deletes", "Net wishlists",
        "Cumulative wishlists", "Email signups (confirmed)", "Cumulative email", "KS pre-launch followers (total)",
        "Discord joins", "Playtest sessions", "Unique playtesters", "Groups that played", "Groups returning another day",
        "Paid spend (EUR)", "Wishlists per 1k visits", "Group return rate", "Cost per net wishlist", "Notes"]
header(ws, 1, cols, [7, 12, 11, 10, 10, 10, 12, 12, 11, 13, 10, 10, 11, 10, 12, 10, 12, 10, 11, 40])
NW = 30
for r in range(2, NW + 2):
    ex = r == 2
    fnt = EX if ex else BLUE
    put(ws, f"A{r}", r - 1, BLACK)
    if ex:
        for col, v in zip("BCDEHJKLMNOP", ["2026-11-02", 640, 38, 3, 21, 0, 15, 2, 9, 3, 1, 0]):
            put(ws, f"{col}{r}", v, fnt)
        put(ws, f"T{r}", "EXAMPLE ROW — replace with week 1", fnt)
    put(ws, f"F{r}", f'=IF(D{r}="","",D{r}-E{r})', BLACK, NUM)
    put(ws, f"G{r}", f'=IF(F{r}="","",SUM($F$2:F{r}))', BLACK, NUM)
    put(ws, f"I{r}", f'=IF(H{r}="","",SUM($H$2:H{r}))', BLACK, NUM)
    put(ws, f"Q{r}", f'=IF(OR(C{r}="",C{r}=0),"",1000*F{r}/C{r})', BLACK, "0.0")
    put(ws, f"R{r}", f'=IF(OR(N{r}="",N{r}=0),"",O{r}/N{r})', BLACK, PCT)
    put(ws, f"S{r}", f'=IF(OR(F{r}="",F{r}<=0,P{r}=""),"",P{r}/F{r})', BLACK, '€#,##0.00')
ws.freeze_panes = "C2"

ws = wb.create_sheet("Clips")
cols = ["Date posted", "Platform", "Concept (A/B/C)", "Video ID (V1–V5, 08 list)", "Hook (first 1.5 s)", "Views (72 h)",
        "Avg watch %", "Link / profile clicks", "Steam visits (UTM)", "Wishlists (UTM)", "Wishlists per 1k visits", "Notes"]
header(ws, 1, cols, [12, 11, 10, 12, 40, 10, 10, 11, 11, 11, 12, 30])
for r in range(2, 202):
    ex = r == 2
    if ex:
        for col, v in zip("ABCDEFGHIJ", ["2026-11-03", "TikTok", "A", "V1", "We had ONE rule: don't talk.", 4200, 0.41, 60, 48, 3]):
            put(ws, f"{col}{r}", v, EX)
        put(ws, f"L{r}", "EXAMPLE ROW", EX)
    put(ws, f"K{r}", f'=IF(OR(I{r}="",I{r}=0),"",1000*J{r}/I{r})', BLACK, "0.0")
ws.freeze_panes = "A2"

ws = wb.create_sheet("Concepts")
header(ws, 1, ["Concept", "Name", "Steam visits (UTM)", "Wishlists (UTM)", "Wishlists per 1k visits", "Enough data? (≥ min visits)", "Leader?"], [10, 34, 16, 16, 18, 22, 10])
put(ws, "A6", "Minimum visits per concept", BOLD)
put(ws, "C6", 300, BLUE, NUM, KEY, "02 §3 decision rule")
for i, (k, n) in enumerate([("A", "It can hear you"), ("B", "Be the shark"), ("C", "Hide in plain sight")], start=2):
    put(ws, f"A{i}", k, BOLD)
    put(ws, f"B{i}", n)
    put(ws, f"C{i}", f'=SUMIFS(Clips!$I:$I,Clips!$C:$C,A{i})', BLACK, NUM)
    put(ws, f"D{i}", f'=SUMIFS(Clips!$J:$J,Clips!$C:$C,A{i})', BLACK, NUM)
    put(ws, f"E{i}", f'=IF(C{i}=0,"",1000*D{i}/C{i})', BLACK, "0.0")
    put(ws, f"F{i}", f'=IF(C{i}>=$C$6,"Yes","Not yet")')
    put(ws, f"G{i}", f'=IF(AND(F{i}="Yes",E{i}<>"",E{i}=MAX($E$2:$E$4)),"◀ leader","")')
put(ws, "A8", "The example row in Clips counts toward concept A until you overwrite it.", EX)

ws = wb.create_sheet("Playtest Survey")
cols = ["Respondent ID", "Date", "Rounds played", "Played shark? (Y/N)", "Fun 1–5", "Would buy at €12.99? (Yes/Maybe/No)",
        "Usual group size", "Would back on Kickstarter? (Yes/Maybe/No)", "Used a mic? (Y/N)", "Top complaint", "Best moment (quote)", "Counts (3+ rounds)?"]
header(ws, 1, cols, [12, 11, 9, 9, 8, 14, 9, 14, 8, 36, 36, 10])
for r in range(2, 502):
    if r == 2:
        for col, v in zip("ABCDEFGHIJK", ["P001", "2026-11-28", 4, "Y", 4, "Yes", 4, "Maybe", "Y", "Shark too fast at the escape", "\"It heard me sneeze\""]):
            put(ws, f"{col}{r}", v, EX)
    put(ws, f"L{r}", f'=IF(C{r}="","",IF(C{r}>=3,1,0))', BLACK, "0")
ws.freeze_panes = "A2"
put(ws, "N1", "Question bank (ask in this order, after round 3)", BOLD)
qs = ["1. How many rounds did you play? (number)", "2. Did you play as the shark at least once?",
      "3. How fun was it, 1–5?", "4. If Reef Run cost €12.99 on Steam, would you buy it? Yes / Maybe / No",
      "5. How many people are usually in your game group?", "6. Would you back it on Kickstarter to get in early? Yes / Maybe / No",
      "7. Did you use a microphone?", "8. What was the most frustrating thing?", "9. What moment will you tell a friend about?"]
for i, q in enumerate(qs, start=2):
    put(ws, f"N{i}", q)
ws.column_dimensions["N"].width = 70

ws = wb.create_sheet("Gates")
header(ws, 1, ["Gate", "Measure", "Current value", "Threshold", "Status", "Source"], [8, 52, 14, 12, 10, 50])
gates = [
    ("G0", "Kickstarter identity + NL bank verification passed (type Yes when done)", "No", "Yes", "text", "01 §3.1, K3 (creator resident in NL)"),
    ("G1", "Steam page live + key distribution confirmed (Yes/No)", "No", "Yes", "text", "S3"),
    ("G2", "Cumulative net wishlists", "=MAX('Weekly KPIs'!G2:G31)", 7500, "num", "02 §9"),
    ("G3", "Expected backers from direct list = followers×0.20 + email×0.04", "=MAX('Weekly KPIs'!J2:J31)*0.2+MAX('Weekly KPIs'!I2:I31)*0.04", 224, "num", "03 Backer Funnel (MVD v2; ≈1,120 KS followers at 20%)"),
    ("G4a", "Share 'Yes' at €12.99 among respondents with 3+ rounds",
     "=IF(SUM('Playtest Survey'!L2:L501)=0,0,COUNTIFS('Playtest Survey'!C2:C501,\">=3\",'Playtest Survey'!F2:F501,\"Yes\")/SUM('Playtest Survey'!L2:L501))", 0.40, "pct", "02 §7 H3"),
    ("G4b", "Group return rate (all weeks)", "=IF(SUM('Weekly KPIs'!N2:N31)=0,0,SUM('Weekly KPIs'!O2:O31)/SUM('Weekly KPIs'!N2:N31))", 0.30, "pct", "02 §7 H2"),
    ("G4c", "Open critical bugs (type the count)", 1, 0, "max", "00 §4"),
    ("G5", "Team hours confirmed through Q4 2027 (Yes/No)", "No", "Yes", "text", "03"),
]
for i, (g, m, cur, thr, kind, src) in enumerate(gates, start=2):
    put(ws, f"A{i}", g, BOLD)
    put(ws, f"B{i}", m)
    isformula = isinstance(cur, str) and cur.startswith("=")
    put(ws, f"C{i}", cur, BLACK if isformula else BLUE, PCT if kind == "pct" else (NUM if kind in ("num", "max") else None))
    put(ws, f"D{i}", thr, BLUE, PCT if kind == "pct" else (NUM if kind in ("num", "max") else None), KEY)
    if kind == "text":
        put(ws, f"E{i}", f'=IF(C{i}=D{i},"PASS","FAIL")')
    elif kind == "max":
        put(ws, f"E{i}", f'=IF(C{i}<=D{i},"PASS","FAIL")')
    else:
        put(ws, f"E{i}", f'=IF(C{i}>=D{i},"PASS","FAIL")')
    put(ws, f"F{i}", src)
last = len(gates) + 1
put(ws, f"B{last+2}", "ALL GATES", BOLD)
put(ws, f"E{last+2}", f'=IF(COUNTIF(E2:E{last},"PASS")={len(gates)},"GO","NO-GO")', BOLD, None, OUTF)
ws.conditional_formatting.add(f"E2:E{last+2}", CellIsRule(operator="equal", formula=['"PASS"'], fill=GOOD))
ws.conditional_formatting.add(f"E2:E{last+2}", CellIsRule(operator="equal", formula=['"GO"'], fill=GOOD))
ws.conditional_formatting.add(f"E2:E{last+2}", CellIsRule(operator="equal", formula=['"FAIL"'], fill=BAD))
ws.conditional_formatting.add(f"E2:E{last+2}", CellIsRule(operator="equal", formula=['"NO-GO"'], fill=BAD))
wb.calculation.fullCalcOnLoad = True
wb.save("assets/data/REEFRUN_validation_tracker.xlsx")
print("saved validation tracker")

# =====================================================================================================
# 2. Campaign dashboard
# =====================================================================================================
wb = Workbook()
readme(wb, [
    ("REEF RUN — Kickstarter campaign dashboard (11_CAMPAIGN_DASHBOARD.md)", TITLE),
    ("Inputs: copy the goal and planned average pledge from 03_funding_model.xlsx (Scenarios), set the launch date. No cross-file links, on purpose.", BLACK),
    ("Pre-launch: one row per day for the 6 weeks before launch (Kickstarter followers from the pre-launch page, email list, wishlists).", BLACK),
    ("Daily: fill one row per campaign day from the Kickstarter creator dashboard (pledged, backers, cancellations, visits by referrer) and your ad accounts.", BLACK),
    ("Forecast: three cumulative-share curves (yellow, assumptions). Forecast = pledged so far ÷ share of the total a campaign usually has by that day.", BLACK),
    ("Triggers: each row turns red when its condition is met. The response for each is in 11_CAMPAIGN_DASHBOARD.md §4.", BOLD),
    ("Grey italic first row in each log = EXAMPLE. Overwrite it with day 1.", BLACK),
])
ws = wb.create_sheet("Inputs")
ws.column_dimensions["A"].width = 52
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 60
inp = [("Funding goal (EUR)", 14000, USD, "From 03_funding_model.xlsx → Scenarios, FUNDING GOAL (MVD)"),
       ("Planned average pledge (EUR)", 37.56, '€#,##0.00', "From 03 → Rewards (blended)"),
       ("Launch date", "2027-04-13", DATE, "Provisional (03 §5). Type a real date."),
       ("Campaign length (days)", 30, NUM, "Max 30 because of Steam key rules (S3)"),
       ("Marketing budget for the campaign (EUR)", 1500, USD, "03 Costs, MVD (live-campaign ads; v2)"),
       ("Net revenue per backer used to judge ads (EUR)", 28.8, '€#,##0.00', "≈ average pledge × (1 − 23.3% deductions incl. VAT), from 03 v2"),
       ("Kickstarter followers at launch (type on launch day)", 0, NUM, "From the pre-launch page"),
       ("Cancellation alert threshold (% of pledged)", 0.05, PCT, "Assumption")]
for i, (lab, v, fmt, note) in enumerate(inp, start=1):
    put(ws, f"A{i}", lab)
    put(ws, f"B{i}", v, BLUE, fmt, KEY if i in (2, 8) else None)
    put(ws, f"C{i}", note)
ws["B3"].value = None
from datetime import date  # noqa: E402
ws["B3"].value = date(2027, 4, 13)
ws["B3"].number_format = DATE

ws = wb.create_sheet("Pre-launch")
header(ws, 1, ["Days to launch", "Date", "KS followers (total)", "Email subscribers (total)", "Steam wishlists (total)",
               "Expected launch backers (followers×20% + email×4%)", "Share of goal covered at launch", "Notes"], [10, 12, 12, 12, 12, 18, 14, 40])
for k, r in enumerate(range(2, 44)):
    put(ws, f"A{r}", 42 - k)
    put(ws, f"B{r}", f"=Inputs!$B$3-A{r}", BLACK, DATE)
    if r == 2:
        for col, v in zip("CDE", [310, 900, 5200]):
            put(ws, f"{col}{r}", v, EX)
        put(ws, f"H{r}", "EXAMPLE ROW", EX)
    put(ws, f"F{r}", f'=IF(C{r}="","",C{r}*0.2+D{r}*0.04)', BLACK, NUM)
    put(ws, f"G{r}", f'=IF(F{r}="","",F{r}*Inputs!$B$2/Inputs!$B$1)', BLACK, PCT)
ws.freeze_panes = "C2"

ws = wb.create_sheet("Daily")
cols = ["Day", "Date", "Pledged today (EUR)", "New backers today", "Cancelled today (EUR)", "Cancelled backers today",
        "Net pledged (cumulative)", "Backers (cumulative)", "Funded %", "Average pledge", "Visits today",
        "Conversion (new backers ÷ visits)", "Visits: Kickstarter", "Visits: email", "Visits: social", "Visits: creators/press",
        "Visits: ads", "Ad spend today (EUR)", "Ad-attributed backers", "Ad cost per backer", "Days left",
        "Needed per remaining day (EUR)", "Notes / what we did"]
header(ws, 1, cols, [5, 11, 11, 10, 11, 10, 12, 11, 9, 10, 9, 11, 10, 9, 9, 10, 9, 10, 10, 10, 8, 12, 40])
for d in range(1, 31):
    r = d + 1
    put(ws, f"A{r}", d)
    put(ws, f"B{r}", f"=Inputs!$B$3+A{r}-1", BLACK, DATE)
    if d == 1:
        for col, v in zip("CDEFKMNOPQRS", [2900, 101, 0, 0, 1650, 700, 380, 320, 150, 100, 60, 2]):
            put(ws, f"{col}{r}", v, EX)
        put(ws, f"W{r}", "EXAMPLE ROW — replace with real day 1", EX)
    put(ws, f"G{r}", f'=IF(C{r}="","",SUM($C$2:C{r})-SUM($E$2:E{r}))', BLACK, USD)
    put(ws, f"H{r}", f'=IF(D{r}="","",SUM($D$2:D{r})-SUM($F$2:F{r}))', BLACK, NUM)
    put(ws, f"I{r}", f'=IF(G{r}="","",G{r}/Inputs!$B$1)', BLACK, PCT)
    put(ws, f"J{r}", f'=IF(OR(H{r}="",H{r}=0),"",G{r}/H{r})', BLACK, '€#,##0.00')
    put(ws, f"L{r}", f'=IF(OR(K{r}="",K{r}=0),"",D{r}/K{r})', BLACK, PCT)
    put(ws, f"T{r}", f'=IF(OR(S{r}="",S{r}=0),"",R{r}/S{r})', BLACK, '€#,##0.00')
    put(ws, f"U{r}", f"=Inputs!$B$4-A{r}", BLACK, NUM)
    put(ws, f"V{r}", f'=IF(OR(G{r}="",U{r}<=0),"",MAX(0,Inputs!$B$1-G{r})/U{r})', BLACK, USD)
ws.freeze_panes = "C2"

ws = wb.create_sheet("Tiers")
header(ws, 1, ["Tier", "Price", "Planned mix", "Backers (type from KS dashboard)", "Pledged", "Actual mix", "vs plan (points)", "Limit", "Remaining"], [22, 9, 11, 14, 11, 10, 11, 8, 10])
tiers = [("Ripple", 5, 0.07, None), ("Early Diver", 12, 0.18, 300), ("Diver", 15, 0.20, None), ("Founding Diver", 29, 0.22, None),
         ("Pod of Four", 49, 0.20, None), ("Founding Pod", 99, 0.10, None), ("Name the Trawler", 199, 0.03, 25), ("No reward / custom", 1, 0.0, None)]
for i, (n, p, m, lim) in enumerate(tiers, start=2):
    put(ws, f"A{i}", n)
    put(ws, f"B{i}", p, BLUE, USD)
    put(ws, f"C{i}", m, BLUE, PCT)
    put(ws, f"D{i}", 0, BLUE, NUM)
    put(ws, f"E{i}", f"=B{i}*D{i}", BLACK, USD, comment="Custom amounts above the tier price aren't captured; use the KS dashboard total for money.")
    put(ws, f"F{i}", f'=IF(SUM($D$2:$D$9)=0,"",D{i}/SUM($D$2:$D$9))', BLACK, PCT)
    put(ws, f"G{i}", f'=IF(F{i}="","",(F{i}-C{i})*100)', BLACK, "0.0")
    if lim:
        put(ws, f"H{i}", lim, BLUE, NUM)
        put(ws, f"I{i}", f"=H{i}-D{i}", BLACK, NUM)
put(ws, "A11", "Totals", BOLD)
put(ws, "D11", "=SUM(D2:D9)", BOLD, NUM)
put(ws, "E11", "=SUM(E2:E9)", BOLD, USD)

ws = wb.create_sheet("Forecast")
header(ws, 1, ["Day", "Conservative: cumulative share of final total by this day", "Baseline", "Optimistic",
               "Forecast final (conservative)", "Forecast final (baseline)", "Forecast final (optimistic)"], [6, 22, 12, 12, 16, 16, 16])
# Cumulative-share curves (ASSUMPTIONS): typical crowdfunding 'U' shape - strong start, flat middle, final push.
# Conservative assumes MORE of the total already arrived early (front-loaded => less to come);
# optimistic assumes a bigger late surge (more to come).


def curve(day, d3, end27):
    if day <= 3:
        return d3 * (day / 3) ** 0.8
    if day <= 27:
        return d3 + (end27 - d3) * (day - 3) / 24
    return end27 + (1 - end27) * ((day - 27) / 3) ** 1.3


for d in range(1, 31):
    r = d + 1
    put(ws, f"A{r}", d)
    for col, (d3, e27) in zip("BCD", [(0.40, 0.85), (0.30, 0.78), (0.22, 0.70)]):
        put(ws, f"{col}{r}", round(curve(d, d3, e27), 4), BLUE, PCT, KEY if d == 3 else None)
    for col, src in zip("EFG", "BCD"):
        put(ws, f"{col}{r}", f'=IF(Daily!G{r}="","",Daily!G{r}/{src}{r})', BLACK, USD)
put(ws, "I1", "Latest forecast", BOLD)
put(ws, "I2", "Last filled day")
put(ws, "J2", '=COUNT(Daily!G2:G31)', BLACK, NUM)
for i, (lab, col) in enumerate([("Conservative", "E"), ("Baseline", "F"), ("Optimistic", "G")], start=3):
    put(ws, f"I{i}", f"{lab} final pledged")
    put(ws, f"J{i}", f'=IF($J$2=0,"",INDEX({col}2:{col}31,$J$2))', BOLD, USD, OUTF)
    put(ws, f"K{i}", f'=IF($J$2=0,"",INDEX({col}2:{col}31,$J$2)/Inputs!$B$1)', BLACK, PCT)
put(ws, "I7", "Curves are assumptions (yellow = the day-3 anchor). Forecasts before day 4 swing wildly; triggers T3/T7 wait for day 4.", EX)
ws.column_dimensions["I"].width = 28
ws.column_dimensions["J"].width = 14

ws = wb.create_sheet("Triggers")
header(ws, 1, ["#", "Trigger", "Value now", "Threshold", "Fired?", "Response (11 §4)"], [5, 60, 12, 12, 9, 50])
trig = [
    ("T1", "Pre-launch: expected launch backers cover < 30% of goal (latest pre-launch row)",
     "=IF(COUNT('Pre-launch'!G2:G43)=0,\"\",INDEX('Pre-launch'!G2:G43,COUNT('Pre-launch'!G2:G43)))", 0.30, "<", "Weak pre-launch: postpone launch (11 §4.1)"),
    ("T2", "Funded % at end of day 2 below 20%", "=IF(Daily!I3=\"\",\"\",Daily!I3)", 0.20, "<", "Slow first 48 h (11 §4.2)"),
    ("T3", "Baseline forecast below 100% of goal (from day 4)", "=IF(OR(Forecast!J2<4,Forecast!K4=\"\"),\"\",Forecast!K4)", 1.0, "<", "Funding below expectations (11 §4.5)"),
    ("T4", "Pledged today < 50% of needed-per-day for 3 days running (latest 3 filled days)",
     "=IF(Forecast!J2<4,\"\",IF(AND(INDEX(Daily!C2:C31,Forecast!J2)<0.5*INDEX(Daily!V2:V31,Forecast!J2-1),INDEX(Daily!C2:C31,Forecast!J2-1)<0.5*INDEX(Daily!V2:V31,Forecast!J2-2),INDEX(Daily!C2:C31,Forecast!J2-2)<0.5*INDEX(Daily!V2:V31,Forecast!J2-3)),1,0))", 1, "=", "Mid-campaign plateau (11 §4.3)"),
    ("T5", "Cancellations above threshold (% of gross pledged)", "=IF(SUM(Daily!C2:C31)=0,\"\",SUM(Daily!E2:E31)/SUM(Daily!C2:C31))", "=Inputs!B8", ">", "Investigate: price, a bad update, card failures"),
    ("T6", "Ad cost per backer above net revenue per backer", "=IF(SUM(Daily!S2:S31)=0,\"\",SUM(Daily!R2:R31)/SUM(Daily!S2:S31))", "=Inputs!B6", ">", "Pause the ads (11 §4)"),
    ("T7", "Optimistic forecast above 300% of goal (from day 4)", "=IF(OR(Forecast!J2<4,Forecast!K5=\"\"),\"\",Forecast!K5)", 3.0, ">", "Sudden attention (11 §4.4)"),
    ("T8", "A limited tier sold out (remaining = 0)", "=IF(OR(Tiers!I3=0,Tiers!I8=0),1,0)", 1, "=", "Say so once in an update; do NOT add another limited tier"),
]
for i, (k, lab, val, thr, op, resp) in enumerate(trig, start=2):
    put(ws, f"A{i}", k, BOLD)
    put(ws, f"B{i}", lab)
    put(ws, f"C{i}", val, BLACK, PCT if k in ("T1", "T2", "T3", "T5", "T7") else ('€#,##0.00' if k == "T6" else "0"))
    put(ws, f"D{i}", thr, BLUE if not str(thr).startswith("=") else BLACK, PCT if k in ("T1", "T2", "T3", "T5", "T7") else ('€#,##0.00' if k == "T6" else "0"), KEY)
    cond = {"<": f"C{i}<D{i}", ">": f"C{i}>D{i}", "=": f"C{i}=D{i}"}[op]
    put(ws, f"E{i}", f'=IF(C{i}="","—",IF({cond},"FIRED","ok"))', BOLD)
    put(ws, f"F{i}", resp)
ws.conditional_formatting.add("E2:E9", CellIsRule(operator="equal", formula=['"FIRED"'], fill=BAD))
ws.conditional_formatting.add("E2:E9", CellIsRule(operator="equal", formula=['"ok"'], fill=GOOD))
wb.calculation.fullCalcOnLoad = True
wb.save("11_campaign_dashboard.xlsx")
print("saved campaign dashboard")
