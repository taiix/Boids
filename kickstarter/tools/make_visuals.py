"""Builds the Kickstarter / social visual set from REEF RUN's own in-engine art.

Run from kickstarter/:  python tools/make_visuals.py
Inputs (all made in-engine by the team, see Marketing/README.md and Marketing/SteamCapsules/README.md):
  ../Marketing/SteamCapsules/master/master_keyart_3840x2160.png   key art without text
  ../Marketing/SteamCapsules/master/title_reefrun_5002x1600.png   REEF RUN logotype (transparent)
  ../Marketing/Screenshots/0*.png                                  6 gameplay screenshots (1920x1080)
  ../Marketing/SteamCapsules/source/fonts + ../Assets/UI Toolkit/Fonts   Barlow family (SIL OFL 1.1)
Funding numbers are read live from 03_funding_model.xlsx so graphics never drift from the model.
Writes into assets/visuals and assets/social, then prints every file with its size.
"""
import math
import os
import sys
import time

from PIL import Image, ImageDraw, ImageFilter, ImageFont

HERE = os.path.dirname(os.path.abspath(__file__))
KS = os.path.dirname(HERE)
ROOT = os.path.dirname(KS)
sys.path.insert(0, HERE)

MK = os.path.join(ROOT, "Marketing")
KEYART = os.path.join(MK, "SteamCapsules", "master", "master_keyart_3840x2160.png")
TITLE = os.path.join(MK, "SteamCapsules", "master", "title_reefrun_5002x1600.png")
if not os.path.exists(TITLE):  # the Marketing copy can be absent (e.g. stashed); a copy lives in the package
    TITLE = os.path.join(KS, "assets", "source", "title_reefrun_5002x1600.png")
SHOTS = {k: os.path.join(MK, "Screenshots", f) for k, f in [
    ("hero", "01_hero_shark_charge.png"), ("catch", "02_action_the_catch.png"), ("hide", "03_depth_hide_in_the_school.png"),
    ("turtle", "04_task_turtle_rescue.png"), ("shadow", "05_world_shadow_overhead.png"), ("trawler", "06_variety_trawler_net.png")]}
FONTS_A = os.path.join(MK, "SteamCapsules", "source", "fonts")
FONTS_B = os.path.join(ROOT, "Assets", "UI Toolkit", "Fonts")
OUT_V = os.path.join(KS, "assets", "visuals")
OUT_S = os.path.join(KS, "assets", "social")
os.makedirs(OUT_V, exist_ok=True)
os.makedirs(OUT_S, exist_ok=True)

# Game UI palette (Assets/UI Toolkit/ReefRunMainMenu.uss / ReefRunLobby.uss)
NAVY = (5, 23, 41)
TEAL = (23, 148, 153)
AQUA = (44, 195, 196)
ORANGE = (255, 184, 77)
TEXT = (238, 246, 247)
DIM = (143, 176, 187)
SHARKRED = (255, 107, 107)
DEEP = (6, 44, 64)  # scrim colour chosen for the Steam capsules (navy turns lavender over cyan water)


def font(name, size):
    for d in (FONTS_A, FONTS_B):
        p = os.path.join(d, name)
        if os.path.exists(p):
            return ImageFont.truetype(p, size)
    raise FileNotFoundError(name)


def H(size):  # headline: the game's logo face
    return font("BarlowCondensed-ExtraBoldItalic.ttf", size)


def HB(size):
    return font("BarlowCondensed-ExtraBold.ttf", size)


def B(size):  # body
    return font("Barlow-SemiBold.ttf", size)


def BR(size):
    try:
        return font("Barlow-Medium.ttf", size)
    except FileNotFoundError:
        return B(size)


def cover(path_or_img, w, h, focus=(0.5, 0.5)):
    im = Image.open(path_or_img).convert("RGB") if isinstance(path_or_img, str) else path_or_img.convert("RGB")
    s = max(w / im.width, h / im.height)
    im = im.resize((math.ceil(im.width * s), math.ceil(im.height * s)), Image.LANCZOS)
    x = int((im.width - w) * focus[0])
    y = int((im.height - h) * focus[1])
    return im.crop((x, y, x + w, y + h))


def shade(im, color=DEEP, alpha=0.55):
    ov = Image.new("RGB", im.size, color)
    return Image.blend(im, ov, alpha)


def vgrad(w, h, color, _unused, a_top, a_bot):
    """A single-colour overlay whose alpha runs linearly from a_top (top edge) to a_bot (bottom edge)."""
    ramp = Image.linear_gradient("L").resize((w, h))  # 0 at top -> 255 at bottom
    alpha = ramp.point(lambda v: int(a_top + (a_bot - a_top) * v / 255))
    g = Image.new("RGBA", (w, h), color + (0,))
    g.putalpha(alpha)
    return g


def title_logo(width):
    t = Image.open(TITLE).convert("RGBA")
    bbox = t.getbbox()
    t = t.crop(bbox)
    s = width / t.width
    return t.resize((int(t.width * s), int(t.height * s)), Image.LANCZOS)


def text_shadow(d_img, xy, txt, f, fill=TEXT, anchor="la", blur=6, offset=(0, 4), alpha=170):
    sh = Image.new("RGBA", d_img.size, (0, 0, 0, 0))
    ImageDraw.Draw(sh).text((xy[0] + offset[0], xy[1] + offset[1]), txt, font=f, fill=(0, 10, 20, alpha), anchor=anchor)
    sh = sh.filter(ImageFilter.GaussianBlur(blur))
    d_img.alpha_composite(sh)
    ImageDraw.Draw(d_img).text(xy, txt, font=f, fill=fill, anchor=anchor)


def wrap(draw, txt, f, maxw):
    words, lines, cur = txt.split(), [], ""
    for w in words:
        t = (cur + " " + w).strip()
        if draw.textlength(t, font=f) <= maxw:
            cur = t
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def rr(draw, box, r, fill=None, outline=None, width=1):
    draw.rounded_rectangle(box, r, fill=fill, outline=outline, width=width)


def arrow(draw, x, y, length=34, color=DIM, width=6):
    """Right-pointing arrow drawn as shapes (the Barlow fonts have no arrow glyph)."""
    draw.line((x, y, x + length - 10, y), fill=color, width=width)
    draw.polygon([(x + length, y), (x + length - 16, y - 12), (x + length - 16, y + 12)], fill=color)


def check(draw, x, y, color, size=26, width=6):
    """A tick mark drawn as lines (no reliable check glyph in the fonts)."""
    draw.line((x, y + size * 0.55, x + size * 0.38, y + size * 0.9, x + size, y + size * 0.1), fill=color, width=width, joint="curve")


written = []


def save(im, folder, name):
    p = os.path.join(folder, name)
    if im.mode == "RGBA" and not name.endswith(".png"):
        im = im.convert("RGB")
    # Something on this machine intermittently locks freshly written images (Errno 22 on overwrite; same trap as
    # Marketing/SteamCapsules/source/compose_capsules.py), so write a temp file and swap it in, with retries.
    tmp = p + ".tmp" + os.path.splitext(name)[1]
    if name.endswith(".jpg"):
        im.convert("RGB").save(tmp, quality=92, optimize=True)
    else:
        im.save(tmp, optimize=True)
    for attempt in range(20):
        try:
            os.replace(tmp, p)
            break
        except OSError:
            time.sleep(0.5)
    else:
        raise OSError(f"could not replace {p}; the new image is at {tmp}")
    written.append((p, im.size))


# ------------------------------------------------------------------ numbers from the model
def model_numbers():
    from xlcalc import compiler
    from openpyxl import load_workbook
    path = os.path.join(KS, "03_funding_model.xlsx")
    xl = compiler(path)
    ws = load_workbook(path)["Scenarios"]
    row = {str(ws.cell(r, 1).value).strip(): r for r in range(2, ws.max_row + 1)}  # rows found by label

    def g(label, col="B"):
        key = next(k for k in row if k.startswith(label))
        return float(xl.evaluate(f"Scenarios!{col}{row[key]}"))

    return {
        "goal": g("FUNDING GOAL"), "stipend": g("Team stipends"), "cash": g("Non-labor cash costs"),
        "cont": g("Contingency"), "budget": g("TOTAL DEVELOPMENT BUDGET"), "unpaid": g("Unpaid founder labor"),
        "hours": g("Remaining labor"), "real_goal": g("FUNDING GOAL", "C"), "backers": g("Backers needed at average"),
    }


N = model_numbers()
N["deduct"] = N["goal"] - N["budget"]


def functional_split():
    """v2 (2026-10-20): the MVD goal split by *what the money does* (how The Boo Croo and Springfield Hotel show it,
    14 §3). Stipends are spread over functions by the MVD hours in the Labor sheet; cash costs by their Costs line;
    contingency and the platform/tax deductions stay separate. Sums exactly to the goal."""
    from openpyxl import load_workbook
    from xlcalc import compiler
    path = os.path.join(KS, "03_funding_model.xlsx")
    wb = load_workbook(path)
    xl = compiler(path)
    groups = {"Programming & netcode": ["Fix client", "Bug backlog", "Performance", "Controller", "Steamworks", "Settings", "Release"],
              "Game design, balance & onboarding": ["Onboarding", "Balance"],
              "Playtests, community & campaign": ["Playtest", "Campaign"],
              "Backer rewards (cosmetics, credits, PDF)": ["Backer rewards", "Making-of"]}
    lab = wb["Labor"]
    hours = {}
    for r in range(2, lab.max_row + 1):
        name = str(lab.cell(r, 1).value or "")
        h = lab.cell(r, 2).value
        if not isinstance(h, (int, float)) or name.startswith("TOTAL"):
            continue
        g = next((k for k, keys in groups.items() if any(name.startswith(x) for x in keys)), None)
        if g:
            hours[g] = hours.get(g, 0) + h
    tot_h = sum(hours.values())
    out = [(k, N["stipend"] * v / tot_h) for k, v in hours.items()]
    cs = wb["Costs"]
    ads = biz = 0.0
    for r in range(2, cs.max_row + 1):
        name = str(cs.cell(r, 1).value or "")
        if not name or name.startswith("TOTAL"):
            continue
        v = float(xl.evaluate(f"Costs!F{r}") or 0)
        if "ads" in name.lower():
            ads += v
        else:
            biz += v
    out += [("Marketing: pre-launch & live ads", ads), ("Steam, KvK, accountant, test hardware", biz),
            ("Contingency (15%)", N["cont"]), ("Platform fees, VAT, tax reserve", N["deduct"])]
    return out


# ------------------------------------------------------------------ 1. project image 1920x1080
def project_image():
    W, Hh = 1920, 1080
    im = cover(KEYART, W, Hh, (0.5, 0.42)).convert("RGBA")
    im.alpha_composite(vgrad(W, Hh, DEEP, DEEP, 0, 120))
    logo = title_logo(int(W * 0.60))
    im.alpha_composite(logo, ((W - logo.width) // 2, int(Hh * 0.50 - logo.height / 2)))
    f = HB(58)
    txt = "3 FISH  ·  1 SHARK  ·  IT CAN HEAR YOU"
    tw = ImageDraw.Draw(im).textlength(txt, font=f)
    y = int(Hh * 0.50 + logo.height / 2 + 34)
    text_shadow(im, ((W - tw) / 2, y), txt, f, fill=TEXT)
    save(im, OUT_V, "ks_project_image_1920x1080.png")
    save(im.resize((1024, 576), Image.LANCZOS), OUT_V, "ks_project_image_1024x576.jpg")


# ------------------------------------------------------------------ 2. story section headers 1360x240 (2x of 680 column)
def headers():
    items = [("the_twist", "THE SHARK CAN HEAR YOU", "hide", (0.5, 0.45)),
             ("how_a_round_works", "HOW A ROUND WORKS", "hero", (0.5, 0.55)),
             ("whats_playable", "WHAT'S PLAYABLE TODAY", "turtle", (0.5, 0.6)),
             ("rewards", "REWARDS", "shadow", (0.5, 0.35)),
             ("where_the_money_goes", "WHERE THE MONEY GOES", "trawler", (0.5, 0.45)),
             ("roadmap", "ROADMAP", "shadow", (0.5, 0.7)),
             ("team", "THE TEAM", "catch", (0.5, 0.5)),
             ("risks", "RISKS & CHALLENGES", "hide", (0.5, 0.8)),
             ("faq", "FAQ", "turtle", (0.5, 0.3)),
             ("use_of_ai", "USE OF AI", "trawler", (0.5, 0.8))]
    W, Hh = 1360, 240
    for slug, label, shot, focus in items:
        im = shade(cover(SHOTS[shot], W, Hh, focus), DEEP, 0.58).convert("RGBA")
        d = ImageDraw.Draw(im)
        f = H(104)
        text_shadow(im, (56, Hh // 2 + 4), label, f, anchor="lm")
        tw = d.textlength(label, font=f)
        d.rectangle((60, Hh // 2 + 58, 60 + min(tw, 420), Hh // 2 + 66), fill=AQUA)
        save(im, OUT_V, f"ks_header_{slug}_1360x240.png")


# ------------------------------------------------------------------ 3. how a round works 1360 x N
def round_infographic():
    steps = [("1", "PICK YOUR POD", "Four friends in a Steam lobby. One of you becomes the great white. The other three are reef fish.", "hero"),
             ("2", "HIDE IN THE SCHOOL", "Swim into a school and you shrink to its size and move with the flock. The shark has to read behaviour, not shapes.", "hide"),
             ("3", "TALK — CAREFULLY", "Co-op tasks need talking: each fish sees a different clue. But the shark sees a ripple wherever someone speaks.", "turtle"),
             ("4", "SURVIVE THE REEF", "Mid-round events change everything: a trawler net sweeps the arena, night falls, a red tide drifts in.", "trawler"),
             ("5", "RUN FOR IT", "When the clock hits zero, a wall lifts and the shark sees everyone. More than half escape: the fish win.", "shadow")]
    W = 1360
    rowh = 330
    Hh = 150 + rowh * len(steps) + 40
    im = Image.new("RGBA", (W, Hh), NAVY + (255,))
    d = ImageDraw.Draw(im)
    text_shadow(im, (W // 2, 80), "ONE ROUND OF REEF RUN · 10 MINUTES", H(70), anchor="mm")
    for i, (n, head, body, shot) in enumerate(steps):
        y = 150 + i * rowh
        thumb = cover(SHOTS[shot], 520, 292).convert("RGBA")
        mask = Image.new("L", thumb.size, 0)
        ImageDraw.Draw(mask).rounded_rectangle((0, 0, thumb.width - 1, thumb.height - 1), 18, fill=255)
        left = i % 2 == 0
        tx = 60 if left else W - 60 - 520
        im.paste(thumb, (tx, y), mask)
        cx = 640 if left else 60
        d.ellipse((cx, y + 18, cx + 76, y + 94), fill=TEAL)
        d.text((cx + 38, y + 56), n, font=HB(56), fill=TEXT, anchor="mm")
        d.text((cx + 96, y + 56), head, font=H(60), fill=AQUA, anchor="lm")
        for j, line in enumerate(wrap(d, body, BR(32), 600)):
            d.text((cx, y + 120 + j * 44), line, font=BR(32), fill=TEXT)
    d.text((W // 2, Hh - 26), "All images are in-game captures (photo mode, no HUD).", font=BR(24), fill=DIM, anchor="mm")
    save(im, OUT_V, "ks_how_a_round_works_1360x1840.png")


# ------------------------------------------------------------------ 4. the twist explainer (illustration, labelled)
def twist_explainer():
    W, Hh = 1360, 820
    im = shade(cover(SHOTS["hide"], W, Hh, (0.5, 0.5)), DEEP, 0.35).convert("RGBA")
    d = ImageDraw.Draw(im)
    panel = Image.new("RGBA", (W, 250), DEEP + (225,))
    im.alpha_composite(panel, (0, Hh - 250))
    # ripple illustration around a point (where a hiding fish talks)
    cx, cy = 520, 330
    for k, (r, a) in enumerate([(40, 255), (85, 190), (135, 120), (190, 60)]):
        ring = Image.new("RGBA", im.size, (0, 0, 0, 0))
        ImageDraw.Draw(ring).ellipse((cx - r, cy - r * 0.62, cx + r, cy + r * 0.62), outline=SHARKRED + (a,), width=6)
        im.alpha_composite(ring)
    d = ImageDraw.Draw(im)
    rr(d, (cx + 40, cy - 150, cx + 440, cy - 70), 16, fill=(5, 23, 41, 230))
    d.text((cx + 60, cy - 110), '"…it\'s right behind you"', font=BR(36), fill=TEXT, anchor="lm")
    cols = [("FISH TALKS", "Proximity voice: only nearby players hear you."),
            ("SHARK SEES A RIPPLE", "Exact up close, up to 8 m off at the edge of hearing."),
            ("2 SECONDS TO RUN", "A clearly heard fish can be locked on. Its HUD warns: SHARK HEARD YOU.")]
    for i, (h, b) in enumerate(cols):
        x = 60 + i * 430
        d.text((x, Hh - 215), h, font=H(44), fill=AQUA if i < 2 else ORANGE)
        for j, line in enumerate(wrap(d, b, BR(28), 380)):
            d.text((x, Hh - 150 + j * 38), line, font=BR(28), fill=TEXT)
        if i < 2:
            arrow(d, x + 372, Hh - 190)
    d.text((W - 24, 24), "Illustration of a game mechanic: ripple drawn for clarity, not the in-game HUD", font=BR(22), fill=DIM, anchor="ra")
    save(im, OUT_V, "ks_the_twist_explainer_1360x820.png")


# ------------------------------------------------------------------ 5. funding allocation donut
def funding_graphic():
    W, Hh = 1360, 780
    im = Image.new("RGBA", (W, Hh), NAVY + (255,))
    d = ImageDraw.Draw(im)
    cols = [AQUA, TEAL, (120, 214, 205), (68, 120, 160), ORANGE, (180, 150, 210), (255, 220, 160), DIM]
    parts = [(lab, v, cols[i % len(cols)]) for i, (lab, v) in enumerate(functional_split()) if v > 0.5]
    total = sum(p[1] for p in parts)
    cx, cy, R, r = 330, 420, 250, 150
    start = -90.0
    for _, v, c in parts:
        ext = 360 * v / total
        d.pieslice((cx - R, cy - R, cx + R, cy + R), start, start + ext, fill=c)
        start += ext
    d.ellipse((cx - r, cy - r, cx + r, cy + r), fill=NAVY)
    d.text((cx, cy - 22), f"€{total:,.0f}", font=HB(70), fill=TEXT, anchor="mm")
    d.text((cx, cy + 34), "funding goal", font=BR(30), fill=DIM, anchor="mm")
    text_shadow(im, (640, 110), "WHERE €" + f"{total:,.0f}" + " GOES", H(64))
    for i, (lab, v, c) in enumerate(parts):
        y = 190 + i * 62
        d.rounded_rectangle((640, y + 6, 668, y + 34), 6, fill=c)
        d.text((688, y), f"€{v:,.0f}", font=HB(36), fill=TEXT)
        d.text((808, y), f"{100 * v / total:.0f}%", font=HB(36), fill=c)
        d.text((888, y + 6), lab, font=BR(24), fill=TEXT)
    note = (f"Not in this chart: about {N['hours']:,.0f} h of remaining work, half of it unpaid founder time worth "
            f"~€{N['unpaid']:,.0f} at market rates. Numbers: 03_funding_model.xlsx.")
    for j, line in enumerate(wrap(d, note, BR(24), W - 120)):
        d.text((60, Hh - 80 + j * 32), line, font=BR(24), fill=DIM)
    save(im, OUT_V, "ks_funding_allocation_1360x780.png")


# ------------------------------------------------------------------ 6. roadmap
def roadmap():
    W, Hh = 1360, 700
    im = shade(cover(SHOTS["shadow"], W, Hh, (0.5, 0.5)), NAVY, 0.78).convert("RGBA")
    d = ImageDraw.Draw(im)
    text_shadow(im, (60, 60), "ROADMAP", H(80))
    ms = [("DONE", "Jun–Oct 2026", "Playable online round: lobby, voice, schools, tasks, events, escape", AQUA, True),
          ("NOW", "Oct 2026–Feb 2027", "Bug fixing, Steam page, public playtests", AQUA, False),
          ("FEB 22–MAR 1", "2027", "Steam Next Fest demo", TEAL, False),
          ("APR–MAY", "2027", "This Kickstarter (30 days)", ORANGE, False),
          ("JULY", "2027", "Backer beta opens", TEAL, False),
          ("Q4", "2027", "Full release on Steam + backer rewards", TEAL, False),
          ("Q1–Q2", "2028", "2nd map, only if €" + f"{N['real_goal']:,.0f}" + " is reached", DIM, False)]
    y0 = 300
    d.line((80, y0, W - 80, y0), fill=DIM, width=4)
    step = (W - 160) / (len(ms) - 1)
    for i, (a, b, c, col, done) in enumerate(ms):
        x = 80 + i * step
        d.ellipse((x - 18, y0 - 18, x + 18, y0 + 18), fill=col if done or col != DIM else NAVY, outline=col, width=5)
        d.text((x, y0 - 70), a, font=HB(34), fill=col, anchor="mm")
        d.text((x, y0 - 36), b, font=BR(22), fill=DIM, anchor="mm")
        for j, line in enumerate(wrap(d, c, BR(22), 170)):
            d.text((x, y0 + 50 + j * 28), line, font=BR(22), fill=TEXT, anchor="ma")
    d.text((60, Hh - 50), "Release window includes a 30% schedule buffer. Every date change is announced in a backer update first.", font=BR(24), fill=DIM)
    save(im, OUT_V, "ks_roadmap_1360x700.png")


# ------------------------------------------------------------------ 7. built vs funded
def built_vs_funded():
    W = 1360
    built = ["4-player online over Steam lobbies", "Proximity voice + 'heard' ripples", "Hide in AI fish schools",
             "Shark: bite, lunge, sonar, grazes", "3 co-op tasks", "9 world events (incl. orcas)", "Escape finale, spectate, ghost voice"]
    funded = ["Online bug fixing + netcode hardening", "Performance + Low/Medium presets", "Onboarding + controls screen",
              "Full controller support + rebinding", "Push-to-talk + colour-blind options", "Steam achievements",
              "Backer rewards in game"]
    Hh = 220 + 64 * max(len(built), len(funded)) + 120
    im = Image.new("RGBA", (W, Hh), NAVY + (255,))
    d = ImageDraw.Draw(im)
    text_shadow(im, (W // 2, 70), "WHAT EXISTS · WHAT YOUR PLEDGE FUNDS", H(62), anchor="mm")
    for col, (head, items, colr, mark) in enumerate([("IN THE BUILD TODAY", built, AQUA, None), ("FUNDED BY THE CAMPAIGN", funded, ORANGE, "+")]):
        x = 60 + col * 650
        rr(d, (x, 140, x + 590, Hh - 110), 22, fill=(9, 35, 56), outline=colr, width=3)
        d.text((x + 34, 172), head, font=HB(40), fill=colr)
        for i, it in enumerate(items):
            y = 240 + i * 64
            if mark:
                d.text((x + 34, y), mark, font=B(32), fill=colr)
            else:
                check(d, x + 34, y + 8, colr)
            for j, line in enumerate(wrap(d, it, BR(28), 500)[:1]):
                d.text((x + 76, y + 2), line, font=BR(28), fill=TEXT)
    d.text((W // 2, Hh - 60), "Not promised at any funding level: bots or <4-player modes, consoles, Mac/Linux.", font=BR(26), fill=DIM, anchor="mm")
    save(im, OUT_V, "ks_built_vs_funded_1360x~.png".replace("~", str(Hh)))


# ------------------------------------------------------------------ 8. reward cards 1200x675
def reward_cards():
    tiers = [("RIPPLE", 5, None, ["Name in Supporters credits", "Backer Discord role", "4K wallpaper pack"], "shadow"),
             ("EARLY DIVER", 12, "LIMITED · 300", ["1 Steam key (beta + full game)", "Supporters credit"], "hide"),
             ("DIVER", 15, None, ["1 Steam key (beta + full game)", "Supporters credit"], "hide"),
             ("POD OF FOUR", 49, "FOR YOUR GROUP", ["4 Steam keys (beta + full game)", "4 names in Supporters credits"], "turtle"),
             ("FOUNDING DIVER", 29, None, ["1 Steam key (beta + full game)", "Founder cosmetics: fish + shark", "Founders credit + Making-of PDF"], "hero"),
             ("FOUNDING POD", 99, None, ["4 Steam keys", "4 Founder cosmetic sets", "4 Founders credits + Making-of PDF"], "catch"),
             ("NAME THE TRAWLER", 199, "LIMITED · 25", ["Founding Diver rewards", "Your name on the trawler's hull", "(moderated, rotates per round)"], "trawler")]
    W, Hh = 1200, 675
    for name, price, badge, lines, shot in tiers:
        im = shade(cover(SHOTS[shot], W, Hh), DEEP, 0.62).convert("RGBA")
        d = ImageDraw.Draw(im)
        d.text((60, 70), name, font=H(92), fill=TEXT)
        d.text((W - 60, 74), f"€{price}", font=HB(96), fill=ORANGE, anchor="ra")
        if badge:
            bw = d.textlength(badge, font=HB(30)) + 40
            rr(d, (60, 190, 60 + bw, 236), 23, fill=TEAL)
            d.text((80, 213), badge, font=HB(30), fill=TEXT, anchor="lm")
        for i, l in enumerate(lines):
            d.text((60, 290 + i * 62), "•  " + l, font=BR(40), fill=TEXT)
        d.text((60, Hh - 60), "Digital only · Steam (Windows) · keys via Kickstarter Pledge Manager", font=BR(24), fill=DIM)
        save(im, OUT_V, f"ks_reward_{name.lower().replace(' ', '_')}_1200x675.png")


# ------------------------------------------------------------------ 9. team template
def team_template():
    W, Hh = 1360, 640
    im = Image.new("RGBA", (W, Hh), NAVY + (255,))
    d = ImageDraw.Draw(im)
    text_shadow(im, (60, 50), "THE TEAM", H(72))
    for i, (handle, role) in enumerate([("[[NAME]] · Stan", "Networking, Steam, UI, VFX, marketing tools"),
                                        ("[[NAME]] · Narniez", "Fish & shark control, flocking, tasks, turtles")]):
        x = 60 + i * 650
        d.ellipse((x, 170, x + 220, 390), outline=AQUA, width=6)
        d.text((x + 110, 280), "[[PHOTO]]", font=HB(30), fill=DIM, anchor="mm")
        d.text((x + 250, 200), handle, font=HB(44), fill=TEXT)
        for j, line in enumerate(wrap(d, role, BR(28), 340)):
            d.text((x + 250, 266 + j * 38), line, font=BR(28), fill=DIM)
    d.text((60, Hh - 110), "Template: replace [[…]] with real names, photos and roles before use. Don't invent credentials.", font=BR(26), fill=ORANGE)
    d.text((60, Hh - 66), "Development since June 2026 · ~100 commits in 16 weeks · Unity HDRP · Steam", font=BR(26), fill=DIM)
    save(im, OUT_V, "ks_team_TEMPLATE_1360x640.png")


# ------------------------------------------------------------------ 10. social set
def social():
    # square announcement
    im = cover(KEYART, 1080, 1080, (0.5, 0.4)).convert("RGBA")
    im.alpha_composite(vgrad(1080, 1080, DEEP, DEEP, 0, 170))
    logo = title_logo(860)
    im.alpha_composite(logo, ((1080 - logo.width) // 2, 560))
    text_shadow(im, (540, 120), "THE SHARK CAN HEAR YOU", H(84), anchor="mm")
    text_shadow(im, (540, 920), "Wishlist on Steam · [[link]]", HB(50), anchor="mm", fill=AQUA)
    save(im, OUT_S, "social_square_announce_1080x1080.png")
    # vertical story / short cover
    im = cover(SHOTS["shadow"], 1080, 1920).convert("RGBA")
    im.alpha_composite(vgrad(1080, 1920, DEEP, DEEP, 60, 210))
    logo = title_logo(900)
    im.alpha_composite(logo, (90, 1250))
    text_shadow(im, (540, 260), "3 FISH", H(150), anchor="mm")
    text_shadow(im, (540, 420), "1 SHARK", H(150), anchor="mm", fill=SHARKRED)
    text_shadow(im, (540, 560), "IT CAN HEAR YOU", H(96), anchor="mm", fill=AQUA)
    text_shadow(im, (540, 1650), "[[CTA: Wishlist / Playtest / Back us]]", HB(52), anchor="mm")
    save(im, OUT_S, "social_vertical_cover_1080x1920.png")
    # X / Bluesky card
    im = cover(SHOTS["hero"], 1600, 900, (0.5, 0.5)).convert("RGBA")
    im.alpha_composite(vgrad(1600, 900, DEEP, DEEP, 0, 200))
    logo = title_logo(640)
    im.alpha_composite(logo, (60, 900 - logo.height - 70))
    text_shadow(im, (1540, 830), "Keep your voice down.", H(70), anchor="rs")
    save(im, OUT_S, "social_card_1600x900.png")
    # YouTube thumbnails: template + one worked example
    for label, hook, shot, col in [("TEMPLATE", "[[HOOK, 3 WORDS]]", "catch", TEXT), ("example_it_heard_us", "IT HEARD US", "catch", SHARKRED)]:
        im = cover(SHOTS[shot], 1280, 720, (0.55, 0.5)).convert("RGBA")
        im.alpha_composite(vgrad(1280, 720, DEEP, DEEP, 140, 0))
        text_shadow(im, (50, 40), hook, H(150), fill=col, blur=10, alpha=230)
        logo = title_logo(360)
        im.alpha_composite(logo, (1280 - logo.width - 36, 720 - logo.height - 30))
        save(im, OUT_S, f"thumbnail_{label}_1280x720.png")


if __name__ == "__main__":
    project_image()
    headers()
    round_infographic()
    twist_explainer()
    funding_graphic()
    roadmap()
    built_vs_funded()
    reward_cards()
    team_template()
    social()
    for p, s in written:
        print(f"{os.path.relpath(p, KS)}  {s[0]}x{s[1]}")
