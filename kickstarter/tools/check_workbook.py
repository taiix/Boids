"""Evaluates every formula in a workbook with pycel (no Excel/LibreOffice needed) and reports errors.

Usage: python tools/check_workbook.py <file.xlsx> [Sheet!A1 ...]
Prints the count of formulas, any cell that evaluates to an Excel error, and the values of the cells listed.
"""
import os
import sys

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(__file__))
from xlcalc import compiler  # noqa: E402

path = sys.argv[1]
wb = load_workbook(path)
xl = compiler(path)
n, errors = 0, []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                n += 1
                addr = f"'{ws.title}'!{c.coordinate}"
                try:
                    v = xl.evaluate(addr)
                except Exception as e:  # noqa: BLE001
                    errors.append((addr, f"EXC {str(e).splitlines()[-1]}"))
                    continue
                if isinstance(v, str) and v.startswith("#"):
                    errors.append((addr, v))
print(f"formulas: {n}  errors: {len(errors)}")
for a, e in errors[:60]:
    print("  ", a, e)
for a in sys.argv[2:]:
    print(a, "=", xl.evaluate(a))
